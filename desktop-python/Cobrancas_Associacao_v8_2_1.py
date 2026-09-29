import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog
import os, sys, json, time, threading, socket, hashlib, secrets, shutil, subprocess, urllib.request, urllib.parse, urllib.error, webbrowser
from datetime import datetime

try:
    import pystray
    from PIL import Image, ImageDraw
except Exception:
    pystray = None

try:
    import pyautogui
except Exception:
    pyautogui = None

try:
    import keyring
except Exception:
    keyring = None

try:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
except Exception:
    SimpleDocTemplate = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except Exception:
    Workbook = None

BASE = os.path.dirname(os.path.abspath(sys.argv[0] if getattr(sys, 'frozen', False) else __file__))
CFG = os.path.join(BASE, 'config.json')
STATE = os.path.join(BASE, 'state.json')
APP = 'CobrancasAssociacaoV82'
VERSION = '8.2.1'
DEFAULT_WEB_APP_URL = 'https://script.google.com/macros/s/AKfycbx4Hm5Nh2J8HeEhctvlRBdKKFDy-QY3miBFz1qZxq_QEHEleF-skUIhfceUL4k8bPxEdg/exec'
LOG_FILE = os.path.join(BASE, 'cobrador.log')
KEYRING_SERVICE = 'CobrancasAssociacao'
INSTANCE_PORT = 47631
INSTANCE_SOCKET = None

def acquire_single_instance():
    global INSTANCE_SOCKET
    try:
        INSTANCE_SOCKET = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        INSTANCE_SOCKET.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 0)
        INSTANCE_SOCKET.bind(('127.0.0.1', INSTANCE_PORT))
        INSTANCE_SOCKET.listen(1)
        return True
    except OSError:
        return False

def password_hash(password, salt):
    return hashlib.pbkdf2_hmac('sha256', password.encode('utf-8'), salt.encode('ascii'), 180000).hex()

def brl_local(v):
    return 'R$ {:,.2f}'.format(float(v or 0)).replace(',','X').replace('.',',').replace('X','.')

def get_saved_key(cfg):
    if keyring:
        try:
            value = keyring.get_password(KEYRING_SERVICE, 'web_app_chave')
            if value: return value
        except Exception: pass
    return str(cfg.get('chave', '') or '')

def store_key(value):
    if keyring:
        try:
            keyring.set_password(KEYRING_SERVICE, 'web_app_chave', value)
            return True
        except Exception: pass
    return False

def load_json(path, default):
    try:
        with open(path, 'r', encoding='utf-8') as f:
            d = default.copy(); d.update(json.load(f)); return d
    except Exception:
        return default.copy()

def save_json(path, data):
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

def http(url, data=None):
    last_error = None
    for attempt in range(2):
        try:
            headers = {'User-Agent': 'CobrancasAssociacaoV82/1.0'}
            if data is None:
                q = urllib.request.Request(url, headers=headers)
            else:
                headers['Content-Type'] = 'text/plain;charset=utf-8'
                q = urllib.request.Request(url, data=json.dumps(data).encode('utf-8'), headers=headers, method='POST')
            r = urllib.request.urlopen(q, timeout=30)
            raw = r.read().decode('utf-8', 'replace'); r.close()
            try: return json.loads(raw)
            except Exception: raise Exception('Web App não retornou JSON válido: ' + raw[:250])
        except urllib.error.HTTPError as e:
            if e.code == 404:
                raise Exception('Web App não encontrado (404). Confira se a URL termina em /exec e se a nova versão foi implantada.')
            last_error = e
        except (urllib.error.URLError, TimeoutError, socket.timeout) as e:
            last_error = e
        if attempt == 0: time.sleep(1.5)
    raise Exception('Falha de conexão com o Web App: ' + str(last_error))

def set_startup(enable, minimized=True):
    if os.name != 'nt': return
    try:
        import winreg
        k = winreg.OpenKey(winreg.HKEY_CURRENT_USER, r'Software\Microsoft\Windows\CurrentVersion\Run', 0, winreg.KEY_SET_VALUE)
        for antigo in ('CobrancasAssociacaoV5','CobrancasAssociacaoV6'):
            try: winreg.DeleteValue(k, antigo)
            except Exception: pass
        suffix = ' --tray' if minimized else ''
        if getattr(sys,'frozen',False): cmd = f'"{sys.executable}"{suffix}'
        else: cmd = f'"{sys.executable.replace("python.exe","pythonw.exe")}" "{os.path.abspath(__file__)}"{suffix}'
        if enable: winreg.SetValueEx(k, APP, 0, winreg.REG_SZ, cmd)
        else:
            try: winreg.DeleteValue(k, APP)
            except Exception: pass
        winreg.CloseKey(k)
    except Exception: pass

class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title('Cobranças da Associação v8.2.1')
        self.geometry('1180x880')
        self.minsize(1040,760)
        self.c = load_json(CFG, {'web_app_url':DEFAULT_WEB_APP_URL,'chave':'','modo':'ASSISTIDO','delay':6,'intervalo':15,'startup':True,'minimizado':False,'password_hash':'','password_salt':'','competencia':'','nome_associacao':'Associação de Moradores','operador':'','pix':''})
        if not self.c.get('web_app_url'): self.c['web_app_url'] = DEFAULT_WEB_APP_URL
        self.c['chave'] = get_saved_key(self.c)
        self.authenticated = self.authenticate()
        if not self.authenticated:
            self.destroy(); return
        self.state = load_json(STATE, {'queue':[], 'index':0, 'retry':[]})
        self.fila = self.state.get('queue', [])
        self.i = int(self.state.get('index',0) or 0)
        self.retry = self.state.get('retry', [])
        self.running = False
        self.stop = False
        self.tray = None
        self.lastsig = None
        self.consulta_cache = None
        self.consulta_em = 0
        self.dashboard_data = None
        self.setup_style()
        set_startup(self.c['startup'], self.c['minimizado'])
        self.build()
        self.protocol('WM_DELETE_WINDOW', self.hide)
        self.make_tray()
        threading.Thread(target=self.monitor, daemon=True).start()
        self.after(1300, self.refresh_dashboard)
        if self.fila and self.i < len(self.fila):
            self.after(1000, self.offer_resume)
        # A abertura manual pelo INICIAR_V8_2.bat deve permanecer visível.
        # O programa só se oculta quando iniciado explicitamente com --tray.
        if '--tray' in sys.argv and self.tray:
            self.after(700, self.hide)

    def authenticate(self):
        expected = str(self.c.get('password_hash','') or '')
        salt = str(self.c.get('password_salt','') or '')
        if not expected or not salt: return True
        self.withdraw()
        for _ in range(3):
            value = simpledialog.askstring('Acesso protegido','Digite a senha do cobrador:',show='*',parent=self)
            if value is None: return False
            if secrets.compare_digest(password_hash(value,salt),expected):
                self.deiconify(); return True
            messagebox.showerror('Senha incorreta','A senha informada está incorreta.',parent=self)
        return False

    def setup_style(self):
        self.configure(bg='#F3F6FA')
        style = ttk.Style(self)
        try: style.theme_use('clam')
        except Exception: pass
        style.configure('.',font=('Segoe UI',10),background='#F3F6FA',foreground='#172033')
        style.configure('TFrame',background='#F3F6FA')
        style.configure('Header.TFrame',background='#123B66')
        style.configure('Header.TLabel',background='#123B66',foreground='white',font=('Segoe UI',22,'bold'))
        style.configure('HeaderSub.TLabel',background='#123B66',foreground='#DCEAF7',font=('Segoe UI',10))
        style.configure('Title.TLabel',background='#F3F6FA',foreground='#123B66',font=('Segoe UI',18,'bold'))
        style.configure('Summary.TLabel',background='#E8F0F8',foreground='#123B66',font=('Segoe UI',11,'bold'),padding=10)
        style.configure('HealthOK.TLabel',background='#DDF4E6',foreground='#176B3A',padding=7)
        style.configure('HealthError.TLabel',background='#FCE3E3',foreground='#A32323',padding=7)
        style.configure('HealthWait.TLabel',background='#FFF3D9',foreground='#8A5A00',padding=7)
        style.configure('Card.TLabelframe',background='white',bordercolor='#D8E1EB',relief='solid')
        style.configure('Card.TLabelframe.Label',background='#F3F6FA',foreground='#123B66',font=('Segoe UI',11,'bold'))
        style.configure('Primary.TButton',background='#1769AA',foreground='white',font=('Segoe UI',10,'bold'),padding=(12,8))
        style.map('Primary.TButton',background=[('active','#0F5791')])
        style.configure('Success.TButton',background='#1E8E5A',foreground='white',font=('Segoe UI',10,'bold'),padding=(12,8))
        style.map('Success.TButton',background=[('active','#167346')])
        style.configure('Warning.TButton',background='#E98B19',foreground='white',font=('Segoe UI',10,'bold'),padding=(12,8))
        style.map('Warning.TButton',background=[('active','#C87109')])
        style.configure('Danger.TButton',background='#C44242',foreground='white',font=('Segoe UI',10,'bold'),padding=(12,8))
        style.map('Danger.TButton',background=[('active','#A73131')])
        style.configure('Light.TButton',background='#E8F0F8',foreground='#123B66',font=('Segoe UI',10,'bold'),padding=(11,7))
        style.map('Light.TButton',background=[('active','#D4E3F2')])
        style.configure('Treeview',rowheight=31,font=('Segoe UI',10),background='white',fieldbackground='white')
        style.configure('Treeview.Heading',background='#123B66',foreground='white',font=('Segoe UI',10,'bold'),padding=7)
        style.map('Treeview.Heading',background=[('active','#174D82')])

    def build(self):
        top = ttk.Frame(self,padding=(20,14),style='Header.TFrame'); top.pack(fill='x')
        titles=ttk.Frame(top,style='Header.TFrame'); titles.pack(side='left')
        ttk.Label(titles,text='Cobranças da Associação',style='Header.TLabel').pack(anchor='w')
        ttk.Label(titles,text='Central financeira e de relacionamento • versão '+VERSION,style='HeaderSub.TLabel').pack(anchor='w')
        ttk.Button(top,text='CONFIGURAÇÃO',style='Light.TButton',command=self.config).pack(side='right')
        ttk.Button(top,text='ATUALIZAR AGORA',style='Light.TButton',command=lambda:self.load(False)).pack(side='right',padx=8)
        status=ttk.Frame(self,padding=(18,10,18,0)); status.pack(fill='x')
        self.health = ttk.Label(status,text='Conexão: verificando...',style='HealthWait.TLabel'); self.health.pack(fill='x')
        self.summary = ttk.Label(self,text='Aguardando atualização...',style='Summary.TLabel'); self.summary.pack(fill='x',padx=18,pady=(6,10))
        cards=ttk.Frame(self,padding=(18,0,18,10)); cards.pack(fill='x')
        self.card_vars={}
        card_defs=[('moradores','MORADORES','#1769AA'),('pagos','PAGOS','#1E8E5A'),('pendentes','A PAGAR','#E98B19'),('aberto','EM ABERTO','#C44242')]
        for i,(key,title,color) in enumerate(card_defs):
            box=tk.Frame(cards,bg='white',highlightbackground='#D8E1EB',highlightthickness=1,padx=14,pady=10)
            box.grid(row=0,column=i,sticky='ew',padx=(0 if i==0 else 5,0))
            tk.Label(box,text=title,bg='white',fg='#627083',font=('Segoe UI',9,'bold')).pack(anchor='w')
            var=tk.StringVar(value='—'); self.card_vars[key]=var
            tk.Label(box,textvariable=var,bg='white',fg=color,font=('Segoe UI',18,'bold')).pack(anchor='w',pady=(3,0))
            cards.columnconfigure(i,weight=1)
        quick = ttk.LabelFrame(self,text='Consultas inteligentes',padding=12,style='Card.TLabelframe'); quick.pack(fill='x',padx=18,pady=(0,10))
        ttk.Button(quick,text='A PAGAR',style='Warning.TButton',command=lambda:self.consultar('pendentes')).pack(side='left')
        ttk.Button(quick,text='PAGOS',style='Success.TButton',command=lambda:self.consultar('pagos')).pack(side='left',padx=7)
        ttk.Button(quick,text='RESUMO',style='Primary.TButton',command=lambda:self.consultar('resumo')).pack(side='left')
        ttk.Button(quick,text='REGISTRAR PAGAMENTO',style='Success.TButton',command=self.registrar_pagamento).pack(side='left',padx=7)
        self.pergunta = tk.StringVar()
        entrada = ttk.Entry(quick,textvariable=self.pergunta,width=38)
        entrada.pack(side='left',padx=(18,5),fill='x',expand=True)
        entrada.insert(0,'Digite: quem está devendo?')
        entrada.bind('<FocusIn>',lambda e:self._limpar_exemplo())
        entrada.bind('<Return>',lambda e:self.perguntar())
        ttk.Button(quick,text='PERGUNTAR',style='Primary.TButton',command=self.perguntar).pack(side='right')
        manage=ttk.LabelFrame(self,text='Gestão da associação',padding=10,style='Card.TLabelframe'); manage.pack(fill='x',padx=18,pady=(0,10))
        ttk.Button(manage,text='MORADORES',style='Primary.TButton',command=self.gerenciar_moradores).pack(side='left')
        ttk.Button(manage,text='MENSAGENS',style='Light.TButton',command=self.editar_mensagens).pack(side='left',padx=6)
        ttk.Button(manage,text='PAINEL ANUAL',style='Light.TButton',command=self.painel_anual).pack(side='left')
        ttk.Button(manage,text='CONFERIR LOTE',style='Warning.TButton',command=self.conferir_lote).pack(side='left',padx=6)
        ttk.Button(manage,text='COMPETÊNCIA',style='Light.TButton',command=self.selecionar_competencia).pack(side='left')
        ttk.Button(manage,text='FECHAR / REABRIR',style='Light.TButton',command=self.alterar_fechamento).pack(side='left',padx=6)
        ttk.Button(manage,text='AVANÇAR MÊS',style='Warning.TButton',command=self.avancar_competencia).pack(side='left')
        ttk.Button(manage,text='COPIAR PIX',style='Success.TButton',command=self.copiar_pix).pack(side='left',padx=6)
        ttk.Button(manage,text='BACKUP AGORA',style='Light.TButton',command=self.backup_online).pack(side='right')
        atual=ttk.LabelFrame(self,text='Cobrança atual',padding=14,style='Card.TLabelframe'); atual.pack(fill='x',padx=18,pady=(0,10))
        self.nome = ttk.Label(atual,text='Nenhuma cobrança carregada',style='Title.TLabel'); self.nome.pack(anchor='w')
        self.info = ttk.Label(atual,text='Selecione Atualizar agora ou Iniciar lote.'); self.info.pack(anchor='w',pady=(4,7))
        self.msg = tk.Text(atual,height=6,wrap='word',font=('Segoe UI',10),bg='#F8FAFC',fg='#172033',relief='solid',borderwidth=1,padx=10,pady=8); self.msg.pack(fill='x',pady=(3,10))
        b = ttk.Frame(atual); b.pack(fill='x')
        ttk.Button(b,text='INICIAR LOTE',style='Primary.TButton',command=self.batch).pack(side='left')
        ttk.Button(b,text='ABRIR WHATSAPP',style='Success.TButton',command=self.openwa).pack(side='left',padx=5)
        ttk.Button(b,text='CONFIRMAR ENVIADO',style='Success.TButton',command=lambda:self.finish('ENVIADO')).pack(side='left',padx=5)
        ttk.Button(b,text='ADIAR 1 DIA',style='Warning.TButton',command=lambda:self.finish('ADIAR_1_DIA')).pack(side='left',padx=5)
        ttk.Button(b,text='PULAR',style='Light.TButton',command=lambda:self.finish('PULADO')).pack(side='left',padx=5)
        ttk.Button(b,text='ERRO / RETENTAR',style='Danger.TButton',command=lambda:self.finish('ERRO')).pack(side='left',padx=5)
        self.prog=ttk.Label(atual,text='0 de 0',font=('Segoe UI',10,'bold')); self.prog.pack(anchor='w',pady=(10,0))
        logf=ttk.LabelFrame(self,text='Atividades',padding=8,style='Card.TLabelframe'); logf.pack(fill='both',expand=True,padx=18,pady=(0,15))
        self.log=tk.Text(logf,height=8,state='disabled',font=('Consolas',9),bg='#101A28',fg='#D8E6F3',insertbackground='white',relief='flat',padx=8,pady=8); self.log.pack(fill='both',expand=True)

    def config(self):
        self.show(); w=tk.Toplevel(self); w.title('Configuração e segurança'); w.geometry('760x680')
        f=ttk.Frame(w,padding=15); f.pack(fill='both',expand=True)
        ttk.Label(f,text='URL do Web App /exec').pack(anchor='w'); u=ttk.Entry(f,width=95); u.pack(fill='x'); u.insert(0,self.c['web_app_url'])
        ttk.Label(f,text='Chave').pack(anchor='w',pady=(10,0)); k=ttk.Entry(f,width=95,show='*'); k.pack(fill='x'); k.insert(0,self.c['chave'])
        dados=ttk.Frame(f);dados.pack(fill='x',pady=(10,0))
        da=ttk.Frame(dados);da.pack(side='left',fill='x',expand=True);db=ttk.Frame(dados);db.pack(side='left',fill='x',expand=True,padx=(10,0))
        ttk.Label(da,text='Nome da associação').pack(anchor='w');assoc=ttk.Entry(da);assoc.pack(fill='x');assoc.insert(0,self.c.get('nome_associacao','Associação de Moradores'))
        ttk.Label(db,text='Responsável / operador').pack(anchor='w');oper=ttk.Entry(db);oper.pack(fill='x');oper.insert(0,self.c.get('operador',''))
        ttk.Label(f,text='Chave PIX padrão').pack(anchor='w',pady=(8,0));pix=ttk.Entry(f);pix.pack(fill='x');pix.insert(0,self.c.get('pix',''))
        row=ttk.Frame(f); row.pack(fill='x',pady=14)
        ttk.Label(row,text='Modo').pack(side='left'); m=ttk.Combobox(row,values=['ASSISTIDO','AUTOMATICO'],state='readonly',width=16); m.pack(side='left',padx=8); m.set(self.c['modo'])
        ttk.Label(row,text='Espera (s)').pack(side='left',padx=(20,0)); d=ttk.Spinbox(row,from_=3,to=30,width=6); d.pack(side='left',padx=8); d.set(self.c['delay'])
        ttk.Label(row,text='Monitor (min)').pack(side='left',padx=(20,0)); it=ttk.Spinbox(row,from_=5,to=120,increment=5,width=6); it.pack(side='left',padx=8); it.set(self.c['intervalo'])
        sv=tk.BooleanVar(value=self.c['startup']); mv=tk.BooleanVar(value=self.c['minimizado'])
        ttk.Checkbutton(f,text='Iniciar automaticamente com o Windows',variable=sv).pack(anchor='w',pady=5)
        ttk.Checkbutton(f,text='Iniciar minimizado na bandeja',variable=mv).pack(anchor='w',pady=5)
        ttk.Label(f,text='No modo ASSISTIDO o sistema nunca marca ENVIADO sozinho. Você confirma com Enter na janela flutuante.',wraplength=700).pack(anchor='w',pady=14)
        security=ttk.LabelFrame(f,text='Segurança e manutenção',padding=10,style='Card.TLabelframe'); security.pack(fill='x',pady=(0,10))
        ttk.Button(security,text='DEFINIR / ALTERAR SENHA',style='Primary.TButton',command=lambda:self.set_password_dialog(w)).pack(side='left')
        ttk.Button(security,text='REMOVER SENHA',style='Light.TButton',command=lambda:self.clear_password(w)).pack(side='left',padx=6)
        ttk.Button(security,text='BACKUP DA CONFIGURAÇÃO',style='Light.TButton',command=self.backup_config).pack(side='left',padx=6)
        ttk.Button(security,text='TESTAR CONEXÃO',style='Success.TButton',command=self.test_connection).pack(side='left')
        ttk.Button(f,text='ATIVAR BACKUP DIÁRIO DA PLANILHA',style='Light.TButton',command=self.ativar_backup_diario).pack(anchor='w')
        def save():
            self.c.update({'web_app_url':u.get().strip() or DEFAULT_WEB_APP_URL,'chave':k.get().strip(),'nome_associacao':assoc.get().strip() or 'Associação de Moradores','operador':oper.get().strip(),'pix':pix.get().strip(),'modo':m.get(),'delay':int(d.get()),'intervalo':int(it.get()),'startup':sv.get(),'minimizado':mv.get()})
            self.save_config_secure(); set_startup(self.c['startup'],self.c['minimizado']); w.destroy(); self.write('Configuração salva com segurança.')
        ttk.Button(f,text='SALVAR CONFIGURAÇÃO',style='Primary.TButton',command=save).pack(anchor='w',pady=10)

    def save_config_secure(self):
        disk = self.c.copy()
        if self.c.get('chave') and store_key(self.c['chave']): disk['chave'] = ''
        save_json(CFG,disk)

    def set_password_dialog(self,parent=None):
        nova=simpledialog.askstring('Nova senha','Digite a nova senha:',show='*',parent=parent or self)
        if not nova: return
        confirma=simpledialog.askstring('Confirmar senha','Digite novamente:',show='*',parent=parent or self)
        if nova != confirma:
            messagebox.showerror('Senha','As senhas não coincidem.',parent=parent or self); return
        if len(nova) < 4:
            messagebox.showwarning('Senha','Use pelo menos 4 caracteres.',parent=parent or self); return
        salt=secrets.token_hex(16); self.c['password_salt']=salt; self.c['password_hash']=password_hash(nova,salt)
        self.save_config_secure(); messagebox.showinfo('Senha','Proteção por senha ativada.',parent=parent or self)

    def clear_password(self,parent=None):
        if not self.c.get('password_hash'):
            messagebox.showinfo('Senha','Nenhuma senha está configurada.',parent=parent or self); return
        if messagebox.askyesno('Remover senha','Deseja remover a senha de acesso?',parent=parent or self):
            self.c['password_hash']=''; self.c['password_salt']=''; self.save_config_secure()
            messagebox.showinfo('Senha','Senha removida.',parent=parent or self)

    def backup_config(self):
        destino=filedialog.asksaveasfilename(title='Salvar backup',defaultextension='.json',filetypes=[('Arquivo JSON','*.json')],initialfile='backup_cobrador_v8_2_1.json')
        if not destino: return
        backup=self.c.copy(); backup['chave']='PROTEGIDA_NO_COMPUTADOR'
        with open(destino,'w',encoding='utf-8') as f: json.dump(backup,f,ensure_ascii=False,indent=2)
        messagebox.showinfo('Backup','Backup salvo com sucesso.')

    def test_connection(self):
        def worker():
            try:
                r=self.api_get('health')
                self.after(0,lambda:messagebox.showinfo('Conexão',f"Conexão OK\nBackend: {r.get('versao','')}\nPlanilha: {r.get('planilha','OK')}"))
            except Exception as e: self.after(0,lambda e=e:messagebox.showerror('Conexão',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def ativar_backup_diario(self):
        def worker():
            try:r=self.api_post({'acao':'ativar_backup'});self.after(0,lambda:messagebox.showinfo('Backup diário',r.get('mensagem','Backup diário ativado.')))
            except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Backup diário',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def api_get(self,acao='fila',**params):
        if self.c.get('competencia') and acao in ('consulta','fila','status'):
            params.setdefault('competencia',self.c['competencia'])
        dados={'acao':acao,'chave':self.c['chave']}; dados.update(params)
        q=urllib.parse.urlencode(dados)
        r=http(self.c['web_app_url']+('&' if '?' in self.c['web_app_url'] else '?')+q)
        if not r.get('ok'): raise Exception(r.get('erro','Falha'))
        return r

    def api_post(self,p):
        p.update({'chave':self.c['chave']})
        r=http(self.c['web_app_url'],p)
        if not r.get('ok'): raise Exception(r.get('erro','Falha'))
        return r

    def _limpar_exemplo(self):
        if self.pergunta.get() == 'Digite: quem está devendo?':
            self.pergunta.set('')

    def copiar_pix(self):
        chave=str(self.c.get('pix','')).strip()
        if not chave:messagebox.showwarning('PIX','Cadastre a chave PIX em Configuração.');return
        self.clipboard_clear();self.clipboard_append(chave);self.update()
        messagebox.showinfo('PIX','Chave PIX copiada para a área de transferência.')

    def perguntar(self):
        q = self.pergunta.get().strip().lower()
        if any(x in q for x in ('devendo','devedor','devedores','pendente','a pagar','não pagou','nao pagou')):
            self.consultar('pendentes')
        elif any(x in q for x in ('pagou','pagaram','pagadores','já pago','ja pago','pagos')):
            self.consultar('pagos')
        elif any(x in q for x in ('resumo','total','quanto','situação','situacao')):
            self.consultar('resumo')
        else:
            messagebox.showinfo('Pergunte ao cobrador','Tente perguntar:\n\n• Quem está devendo?\n• Quem já pagou?\n• Qual é o resumo do mês?')

    def consultar(self,tipo):
        if not self.c.get('web_app_url') or not self.c.get('chave'):
            messagebox.showwarning('Configuração necessária','Informe a URL /exec e a chave em Configuração antes de consultar.')
            return
        self.health.config(text='Conexão: consultando a planilha...',style='HealthWait.TLabel')
        def worker():
            try:
                agora = time.time()
                if self.consulta_cache and agora-self.consulta_em < 30:
                    r = self.consulta_cache
                else:
                    r = self.api_get('consulta')
                    self.consulta_cache = r; self.consulta_em = agora
                self.after(0,lambda:self._mostrar_consulta(tipo,r))
            except Exception as e:
                if 'Ação inválida' in str(e) or 'Acao invalida' in str(e):
                    e = Exception('O Backend v8.2 ainda não foi implantado. Siga o arquivo LEIA-ME.txt e publique uma nova versão do Apps Script.')
                self.after(0,lambda e=e:self._erro_consulta(e))
        threading.Thread(target=worker,daemon=True).start()

    def _erro_consulta(self,e):
        self.health.config(text='Conexão: ERRO — '+str(e),style='HealthError.TLabel')
        messagebox.showerror('Erro na consulta',str(e))

    def _mostrar_consulta(self,tipo,r):
        self.health.config(text='Conexão: OK — Planilha e Apps Script online',style='HealthOK.TLabel')
        self.update_dashboard(r)
        if tipo == 'resumo':
            self._mostrar_resumo(r); return
        itens = r.get(tipo,[])
        titulo = 'Quem já pagou' if tipo == 'pagos' else 'Quem está devendo'
        w=tk.Toplevel(self); w.title(titulo); w.geometry('980x650'); w.transient(self)
        cab=ttk.Frame(w,padding=14); cab.pack(fill='x')
        ttk.Label(cab,text=titulo,font=('Segoe UI',18,'bold')).pack(side='left')
        total = r.get('total_pago','R$ 0,00') if tipo=='pagos' else r.get('total_pendente','R$ 0,00')
        ttk.Label(cab,text=f"{len(itens)} pessoa(s) • {total} • {r.get('competencia','')}",font=('Segoe UI',11,'bold')).pack(side='right')
        filtro=ttk.Frame(w,padding=(14,0,14,8)); filtro.pack(fill='x')
        ttk.Label(filtro,text='Pesquisar:').pack(side='left')
        busca=tk.StringVar(); entrada=ttk.Entry(filtro,textvariable=busca,width=42); entrada.pack(side='left',padx=7)
        ttk.Label(filtro,text='Digite parte do nome, código ou telefone.').pack(side='left')
        if tipo == 'pagos':
            cols=('codigo','nome','telefone','valor','data','forma')
            heads=('Cód.','Morador','Telefone','Valor pago','Data','Forma')
            widths=(55,250,145,100,100,110)
        else:
            cols=('codigo','nome','telefone','valor','vencimento')
            heads=('Cód.','Morador','Telefone','Saldo','Vencimento')
            widths=(55,290,155,110,110)
        area=ttk.Frame(w,padding=(14,0,14,14)); area.pack(fill='both',expand=True)
        tree=ttk.Treeview(area,columns=cols,show='headings',selectmode='browse')
        sy=ttk.Scrollbar(area,orient='vertical',command=tree.yview); tree.configure(yscrollcommand=sy.set)
        for c,h,wd in zip(cols,heads,widths): tree.heading(c,text=h); tree.column(c,width=wd,anchor='w')
        tree.tag_configure('par',background='#F2F6FA')
        tree.tag_configure('impar',background='white')
        visiveis=[]
        def render(*_):
            tree.delete(*tree.get_children()); visiveis.clear(); termo=busca.get().strip().lower()
            for x in itens:
                alvo=f"{x.get('codigo','')} {x.get('morador','')} {x.get('telefone','')}".lower()
                if termo and termo not in alvo: continue
                visiveis.append(x); indice=len(visiveis)-1
                if tipo=='pagos': vals=(x.get('codigo',''),x.get('morador',''),x.get('telefone',''),x.get('valor_pago',''),x.get('data_pagamento',''),x.get('forma_pagamento',''))
                else: vals=(x.get('codigo',''),x.get('morador',''),x.get('telefone',''),x.get('saldo',''),x.get('vencimento',''))
                tree.insert('', 'end',iid=str(indice),values=vals,tags=('par' if indice%2==0 else 'impar',))
        def selecionado():
            sel=tree.selection()
            if not sel: messagebox.showwarning('Seleção','Selecione um morador na lista.',parent=w); return None
            return visiveis[int(sel[0])]
        busca.trace_add('write',render); render()
        tree.pack(side='left',fill='both',expand=True); sy.pack(side='right',fill='y')
        actions=ttk.Frame(w,padding=(14,0,14,14)); actions.pack(fill='x')
        if tipo=='pendentes':
            ttk.Button(actions,text='ABRIR WHATSAPP',style='Success.TButton',command=lambda:self.open_item_whatsapp(selecionado())).pack(side='left')
            ttk.Button(actions,text='REGISTRAR PAGAMENTO',style='Primary.TButton',command=lambda:self.registrar_pagamento(selecionado())).pack(side='left',padx=6)
            tree.bind('<Double-1>',lambda e:self.open_item_whatsapp(selecionado()))
        else:
            ttk.Button(actions,text='ANEXAR COMPROVANTE',style='Primary.TButton',command=lambda:self.anexar_comprovante_pago(selecionado())).pack(side='left')
            ttk.Button(actions,text='EMITIR / ENVIAR RECIBO',style='Success.TButton',command=lambda:self.emitir_recibo_pago(selecionado())).pack(side='left',padx=6)
        ttk.Button(actions,text='HISTÓRICO',style='Light.TButton',command=lambda:self.show_history(selecionado())).pack(side='left',padx=6)
        ttk.Button(actions,text='EXPORTAR PDF',style='Light.TButton',command=lambda:self.export_pdf(tipo,itens,r)).pack(side='right')
        ttk.Button(actions,text='EXPORTAR EXCEL',style='Light.TButton',command=lambda:self.export_excel(tipo,itens,r)).pack(side='right',padx=6)

    def _mostrar_resumo(self,r):
        w=tk.Toplevel(self); w.title('Resumo do mês'); w.geometry('560x390'); w.transient(self)
        ttk.Label(w,text='Resumo do mês',font=('Segoe UI',20,'bold')).pack(pady=(24,4))
        ttk.Label(w,text=r.get('competencia',''),font=('Segoe UI',12)).pack()
        f=ttk.Frame(w,padding=20); f.pack(fill='both',expand=True)
        linhas=[('Moradores',r.get('total_moradores',0)),('Pagos',f"{r.get('qtd_pagos',0)} — {r.get('total_pago','R$ 0,00')}"),('A pagar',f"{r.get('qtd_pendentes',0)} — {r.get('total_pendente','R$ 0,00')}"),('Previsto',r.get('total_previsto','R$ 0,00'))]
        for i,(rot,val) in enumerate(linhas):
            ttk.Label(f,text=rot,font=('Segoe UI',12)).grid(row=i,column=0,sticky='w',pady=10)
            ttk.Label(f,text=str(val),font=('Segoe UI',13,'bold')).grid(row=i,column=1,sticky='e',pady=10,padx=(35,0))
        f.columnconfigure(1,weight=1)
        actions=ttk.Frame(w,padding=(20,0,20,16)); actions.pack(fill='x')
        ttk.Button(actions,text='EXPORTAR PDF',style='Light.TButton',command=lambda:self.export_pdf('resumo',[],r)).pack(side='right')
        ttk.Button(actions,text='EXPORTAR EXCEL',style='Light.TButton',command=lambda:self.export_excel('resumo',[],r)).pack(side='right',padx=6)

    def update_dashboard(self,r):
        self.dashboard_data=r
        self.card_vars['moradores'].set(str(r.get('total_moradores',0)))
        self.card_vars['pagos'].set(str(r.get('qtd_pagos',0)))
        self.card_vars['pendentes'].set(str(r.get('qtd_pendentes',0)))
        self.card_vars['aberto'].set(r.get('total_pendente','R$ 0,00'))
        self.summary.config(text=f"Competência {r.get('competencia','')}  •  Recebido: {r.get('total_pago','R$ 0,00')}  •  Em aberto: {r.get('total_pendente','R$ 0,00')}  •  Previsto: {r.get('total_previsto','R$ 0,00')}")

    def refresh_dashboard(self):
        if not self.c.get('web_app_url') or not self.c.get('chave'): return
        def worker():
            try:
                r=self.api_get('consulta'); self.consulta_cache=r; self.consulta_em=time.time()
                self.after(0,lambda:self.update_dashboard(r))
            except Exception as e: self.write('Painel: '+str(e))
        threading.Thread(target=worker,daemon=True).start()

    def open_item_whatsapp(self,item):
        if not item: return
        n=''.join(c for c in str(item.get('telefone','')) if c.isdigit()); n=n if n.startswith('55') else '55'+n
        mensagem=item.get('mensagem') or f"Olá {item.get('morador','')}, tudo bem? Consta em nosso controle a contribuição da Associação de Moradores referente a {self.consulta_cache.get('competencia','') if self.consulta_cache else ''}, no valor de {item.get('saldo','')}. Caso já tenha realizado o pagamento, por favor desconsidere e nos informe. Obrigado."
        webbrowser.open('https://wa.me/'+n+'?text='+urllib.parse.quote(mensagem))

    def registrar_pagamento(self,item=None):
        if not self.consulta_cache:
            messagebox.showinfo('Atualização','Aguarde o painel carregar ou clique em ATUALIZAR AGORA.'); self.refresh_dashboard(); return
        pendentes=self.consulta_cache.get('pendentes',[])
        if not pendentes:
            messagebox.showinfo('Pagamento','Não existem moradores pendentes neste mês.'); return
        w=tk.Toplevel(self); w.title('Registrar pagamento'); w.geometry('560x510'); w.transient(self); w.grab_set()
        f=ttk.Frame(w,padding=18); f.pack(fill='both',expand=True)
        ttk.Label(f,text='Registrar pagamento',font=('Segoe UI',19,'bold')).pack(anchor='w')
        ttk.Label(f,text='O lançamento será gravado diretamente na mensalidade atual.').pack(anchor='w',pady=(2,14))
        nomes=[f"{x.get('codigo')} — {x.get('morador')}" for x in pendentes]
        ttk.Label(f,text='Morador').pack(anchor='w'); mor=ttk.Combobox(f,values=nomes,state='readonly'); mor.pack(fill='x',pady=(2,10))
        escolhido=0
        if item:
            for i,x in enumerate(pendentes):
                if str(x.get('codigo'))==str(item.get('codigo')): escolhido=i; break
        mor.current(escolhido)
        row=ttk.Frame(f); row.pack(fill='x')
        left=ttk.Frame(row); left.pack(side='left',fill='x',expand=True)
        right=ttk.Frame(row); right.pack(side='left',fill='x',expand=True,padx=(12,0))
        ttk.Label(left,text='Valor pago').pack(anchor='w'); valor=ttk.Entry(left); valor.pack(fill='x'); valor.insert(0,str(pendentes[escolhido].get('saldo','R$ 0,00')).replace('R$','').strip())
        ttk.Label(right,text='Data (dd/mm/aaaa)').pack(anchor='w'); data=ttk.Entry(right); data.pack(fill='x'); data.insert(0,datetime.now().strftime('%d/%m/%Y'))
        ttk.Label(f,text='Forma de pagamento').pack(anchor='w',pady=(12,0)); forma=ttk.Combobox(f,values=['PIX','Dinheiro','Transferencia','Boleto','Cartao','Outro'],state='readonly'); forma.pack(fill='x'); forma.set('PIX')
        ttk.Label(f,text='Observações').pack(anchor='w',pady=(12,0)); obs=tk.Text(f,height=4,wrap='word'); obs.pack(fill='x')
        comprovante=tk.StringVar()
        arq=ttk.Frame(f); arq.pack(fill='x',pady=(8,0))
        ttk.Button(arq,text='ANEXAR COMPROVANTE',style='Light.TButton',command=lambda:comprovante.set(filedialog.askopenfilename(title='Selecionar comprovante',filetypes=[('Comprovantes','*.pdf *.png *.jpg *.jpeg'),('Todos','*.*')]))).pack(side='left')
        ttk.Label(arq,textvariable=comprovante).pack(side='left',padx=8)
        def selecionar(_=None):
            idx=mor.current()
            if idx>=0:
                valor.delete(0,'end'); valor.insert(0,str(pendentes[idx].get('saldo','R$ 0,00')).replace('R$','').strip())
        mor.bind('<<ComboboxSelected>>',selecionar)
        def salvar():
            try:
                idx=mor.current(); x=pendentes[idx]
                numero=float(valor.get().strip().replace('.','').replace(',','.'))
                if numero<=0: raise ValueError('Informe um valor maior que zero.')
                datetime.strptime(data.get().strip(),'%d/%m/%Y')
            except Exception as e:
                messagebox.showerror('Pagamento','Confira o valor e a data.\n'+str(e),parent=w); return
            if not messagebox.askyesno('Confirmar',f"Registrar {valor.get()} para {x.get('morador')}?",parent=w): return
            def worker():
                try:
                    nota=obs.get('1.0','end').strip()
                    if comprovante.get(): nota+=((' — ' if nota else '')+'Comprovante: '+os.path.basename(comprovante.get()))
                    retorno=self.api_post({'acao':'registrar_pagamento','codigo':x.get('codigo'),'valor':numero,'data':data.get().strip(),'forma':forma.get(),'observacao':nota})
                    try: retorno['agradecimento']=self.api_get('templates').get('templates',{}).get('agradecimento','')
                    except Exception: retorno['agradecimento']=''
                    recibo=self.gerar_recibo(x,numero,data.get().strip(),forma.get(),retorno,comprovante.get())
                    self.consulta_cache=None
                    self.after(0,lambda:self.pagamento_concluido(w,x,numero,data.get().strip(),forma.get(),retorno,recibo))
                except Exception as e: self.after(0,lambda e=e:messagebox.showerror('Pagamento',str(e),parent=w))
            threading.Thread(target=worker,daemon=True).start()
        ttk.Button(f,text='REGISTRAR PAGAMENTO',style='Success.TButton',command=salvar).pack(anchor='e',pady=18)

    def pagamento_concluido(self,janela,item,valor,data,forma,retorno,recibo):
        comp=(self.consulta_cache or {}).get('competencia','')
        janela.destroy(); self.refresh_dashboard(); self.load(True)
        texto='Pagamento registrado com sucesso.\n\nRecibo: '+recibo
        if messagebox.askyesno('Pagamento',texto+'\n\nDeseja enviar o recibo pelo WhatsApp?'):
            modelo=retorno.get('agradecimento') or 'Olá [nome], recebemos seu pagamento de [valor] referente a [competencia]. Muito obrigado pela colaboração com a Associação de Moradores!'
            msg=modelo.replace('[nome]',str(item.get('morador',''))).replace('[valor]',brl_local(valor)).replace('[competencia]',comp)
            self.enviar_recibo_whatsapp(item,recibo,msg)

    def show_history(self,item):
        if not item: return
        def worker():
            try:
                r=self.api_get('historico',codigo=item.get('codigo'))
                self.after(0,lambda:self._history_window(item,r.get('historico',[])))
            except Exception as e: self.after(0,lambda e=e:messagebox.showerror('Histórico',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def gerar_recibo(self,item,valor,data,forma,retorno,comprovante=''):
        pasta=os.path.join(BASE,'recibos'); os.makedirs(pasta,exist_ok=True)
        nome='recibo_'+str(item.get('codigo',''))+'_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'.pdf'
        destino=os.path.join(pasta,nome)
        if comprovante:
            try:
                anexos=os.path.join(BASE,'comprovantes'); os.makedirs(anexos,exist_ok=True)
                shutil.copy2(comprovante,os.path.join(anexos,datetime.now().strftime('%Y%m%d_%H%M%S_')+os.path.basename(comprovante)))
            except Exception as e: self.write('Não foi possível copiar o comprovante: '+str(e))
        if SimpleDocTemplate is None: return 'componente PDF não instalado'
        styles=getSampleStyleSheet(); story=[Paragraph('RECIBO DE CONTRIBUIÇÃO',styles['Title']),Paragraph(str(self.c.get('nome_associacao') or 'Associação de Moradores'),styles['Heading2']),Spacer(1,16)]
        dados=[['Morador',item.get('morador','')],['Competência',(self.consulta_cache or {}).get('competencia','')],['Valor recebido',brl_local(valor)],['Data',data],['Forma',forma],['Situação',retorno.get('status','')],['Saldo restante',retorno.get('saldo','')]]
        tab=Table(dados,colWidths=[120,340]); tab.setStyle(TableStyle([('BACKGROUND',(0,0),(0,-1),colors.HexColor('#E8F0F8')),('FONTNAME',(0,0),(0,-1),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),.5,colors.HexColor('#CBD5E1')),('PADDING',(0,0),(-1,-1),9)])); responsavel=str(self.c.get('operador') or 'Responsável pelo recebimento'); story.extend([tab,Spacer(1,22),Paragraph('Recebemos o valor acima referente à contribuição da Associação de Moradores.',styles['BodyText']),Spacer(1,40),Paragraph('____________________________________<br/>'+responsavel,styles['BodyText'])])
        SimpleDocTemplate(destino,pagesize=A4,rightMargin=42,leftMargin=42,topMargin=42,bottomMargin=42).build(story)
        return destino

    def emitir_recibo_pago(self,item):
        if not item:return
        try: valor=float(str(item.get('valor_pago','0')).replace('R$','').replace('.','').replace(',','.').strip())
        except Exception: valor=0
        retorno={'status':'PAGO','saldo':'R$ 0,00'}
        destino=self.gerar_recibo(item,valor,item.get('data_pagamento',''),item.get('forma_pagamento',''),retorno)
        if destino.startswith('componente'):
            messagebox.showerror('Recibo',destino);return
        if messagebox.askyesno('Recibo','Recibo emitido com sucesso:\n'+destino+'\n\nDeseja enviá-lo pelo WhatsApp?'):
            comp=(self.consulta_cache or {}).get('competencia','')
            msg=f"Olá {item.get('morador','')}, segue o recibo de pagamento da contribuição referente a {comp}. Muito obrigado pela colaboração!"
            self.enviar_recibo_whatsapp(item,destino,msg)

    def enviar_recibo_whatsapp(self,item,recibo,mensagem):
        if not recibo or not os.path.isfile(recibo):
            messagebox.showerror('WhatsApp','O arquivo do recibo não foi encontrado.');return
        n=''.join(c for c in str(item.get('telefone','')) if c.isdigit()); n=n if n.startswith('55') else '55'+n
        webbrowser.open('https://wa.me/'+n+'?text='+urllib.parse.quote(mensagem))
        try:
            if os.name=='nt': subprocess.Popen(['explorer.exe','/select,'+os.path.normpath(recibo)])
            else: webbrowser.open('file:///'+os.path.dirname(recibo))
        except Exception: pass
        messagebox.showinfo('Enviar recibo','A conversa foi aberta com a mensagem pronta e o recibo está selecionado na pasta.\n\nArraste o PDF para a conversa ou use o clipe do WhatsApp e confirme o envio.')

    def anexar_comprovante_pago(self,item):
        if not item:return
        origem=filedialog.askopenfilename(title='Selecionar comprovante de '+str(item.get('morador','')),filetypes=[('Comprovantes','*.pdf *.png *.jpg *.jpeg'),('Todos os arquivos','*.*')])
        if not origem:return
        obs=simpledialog.askstring('Comprovante','Observação opcional:',parent=self) or ''
        pasta=os.path.join(BASE,'comprovantes');os.makedirs(pasta,exist_ok=True)
        nome=datetime.now().strftime('%Y%m%d_%H%M%S_')+str(item.get('codigo',''))+'_'+os.path.basename(origem)
        destino=os.path.join(pasta,nome)
        try:
            shutil.copy2(origem,destino)
            self.api_post({'acao':'anexar_comprovante','codigo':item.get('codigo'),'morador':item.get('morador'),'arquivo':nome,'observacao':obs})
            messagebox.showinfo('Comprovante','Comprovante anexado e registrado no histórico.\n\n'+destino)
        except Exception as e:
            messagebox.showerror('Comprovante',str(e))

    def conferir_lote(self):
        if not self.fila:
            if self.load(False)<=0: return
        w=tk.Toplevel(self); w.title('Conferência antes do envio'); w.geometry('1050x620'); w.transient(self)
        ttk.Label(w,text=f'Conferência do lote — {len(self.fila)} mensagem(ns)',font=('Segoe UI',18,'bold'),padding=14).pack(anchor='w')
        cols=('nome','telefone','acao','valor','mensagem'); heads=('Morador','Telefone','Ação','Valor','Mensagem'); widths=(210,125,120,90,430)
        tree=ttk.Treeview(w,columns=cols,show='headings')
        for c,h,wd in zip(cols,heads,widths): tree.heading(c,text=h); tree.column(c,width=wd,anchor='w')
        for x in self.fila: tree.insert('','end',values=(x.get('morador',''),x.get('telefone',''),x.get('acao',''),x.get('saldo',''),x.get('mensagem','')))
        tree.pack(fill='both',expand=True,padx=14,pady=(0,12))
        ttk.Button(w,text='LOTE CONFERIDO — INICIAR',style='Success.TButton',command=lambda:(w.destroy(),self.batch())).pack(pady=(0,14))

    def gerenciar_moradores(self):
        def abrir(r):
            moradores=r.get('moradores',[]); w=tk.Toplevel(self); w.title('Cadastro de moradores'); w.geometry('900x620'); w.transient(self)
            ttk.Label(w,text='Cadastro de moradores',font=('Segoe UI',19,'bold'),padding=14).pack(anchor='w')
            tree=ttk.Treeview(w,columns=('codigo','nome','unidade','telefone','situacao'),show='headings');
            for c,h,wd in [('codigo','Cód.',60),('nome','Nome',280),('unidade','Unidade',150),('telefone','WhatsApp',160),('situacao','Situação',100)]: tree.heading(c,text=h); tree.column(c,width=wd,anchor='w')
            for i,x in enumerate(moradores): tree.insert('','end',iid=str(i),values=(x.get('codigo'),x.get('nome'),x.get('unidade',''),x.get('telefone'),x.get('situacao')))
            tree.pack(fill='both',expand=True,padx=14)
            def editar(novo=False):
                item={} if novo else (moradores[int(tree.selection()[0])] if tree.selection() else None)
                if item is None: messagebox.showwarning('Cadastro','Selecione um morador.',parent=w); return
                d=tk.Toplevel(w); d.title('Morador'); d.geometry('500x350'); d.transient(w); d.grab_set(); f=ttk.Frame(d,padding=16); f.pack(fill='both',expand=True)
                campos={}
                for rot,key in [('Código','codigo'),('Nome completo','nome'),('Unidade / endereço','unidade'),('WhatsApp','telefone')]: ttk.Label(f,text=rot).pack(anchor='w'); e=ttk.Entry(f); e.pack(fill='x',pady=(2,9)); e.insert(0,str(item.get(key,''))); campos[key]=e
                ttk.Label(f,text='Situação').pack(anchor='w'); st=ttk.Combobox(f,values=['Ativo','Inativo'],state='readonly'); st.pack(fill='x'); st.set(item.get('situacao','Ativo'))
                def salvar():
                    p={'acao':'salvar_morador','linha':item.get('linha',''),'codigo':campos['codigo'].get(),'nome':campos['nome'].get(),'unidade':campos['unidade'].get(),'telefone':campos['telefone'].get(),'situacao':st.get()}
                    try:self.api_post(p);d.destroy();w.destroy();messagebox.showinfo('Cadastro','Morador salvo. As próximas mensalidades usarão o cadastro atualizado.');self.refresh_dashboard()
                    except Exception as e:messagebox.showerror('Cadastro',str(e),parent=d)
                ttk.Button(f,text='SALVAR',style='Success.TButton',command=salvar).pack(anchor='e',pady=16)
            b=ttk.Frame(w,padding=14); b.pack(fill='x'); ttk.Button(b,text='NOVO MORADOR',style='Success.TButton',command=lambda:editar(True)).pack(side='left'); ttk.Button(b,text='EDITAR SELECIONADO',style='Primary.TButton',command=editar).pack(side='left',padx=6)
        def worker():
            try:r=self.api_get('moradores');self.after(0,lambda:abrir(r))
            except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Moradores',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def editar_mensagens(self):
        def abrir(r):
            w=tk.Toplevel(self);w.title('Mensagens automáticas');w.geometry('850x700');w.transient(self);f=ttk.Frame(w,padding=14);f.pack(fill='both',expand=True)
            ttk.Label(f,text='Mensagens automáticas',font=('Segoe UI',19,'bold')).pack(anchor='w');ttk.Label(f,text='Campos disponíveis: [nome], [competencia], [valor] e [vencimento]').pack(anchor='w',pady=(2,10))
            caixas={}; nomes=[('lembrete','Lembrete antecipado'),('vence_hoje','Vence hoje'),('primeira','Primeira cobrança'),('segunda','Segunda cobrança'),('agradecimento','Agradecimento')]
            for k,rot in nomes: ttk.Label(f,text=rot,font=('Segoe UI',10,'bold')).pack(anchor='w');t=tk.Text(f,height=4,wrap='word');t.pack(fill='x',pady=(2,8));t.insert('1.0',r.get('templates',{}).get(k,''));caixas[k]=t
            def salvar():
                try:self.api_post({'acao':'salvar_templates','templates':{k:t.get('1.0','end').strip() for k,t in caixas.items()}});w.destroy();messagebox.showinfo('Mensagens','Mensagens atualizadas.')
                except Exception as e:messagebox.showerror('Mensagens',str(e),parent=w)
            ttk.Button(f,text='SALVAR MENSAGENS',style='Success.TButton',command=salvar).pack(anchor='e')
        threading.Thread(target=lambda:self._thread_ui('templates',abrir,'Mensagens'),daemon=True).start()

    def _thread_ui(self,acao,callback,titulo,**params):
        try:r=self.api_get(acao,**params);self.after(0,lambda:callback(r))
        except Exception as e:self.after(0,lambda e=e:messagebox.showerror(titulo,str(e)))

    def selecionar_competencia(self):
        atual=self.c.get('competencia') or (self.consulta_cache or {}).get('competencia','')
        valor=simpledialog.askstring('Competência','Informe MM/AAAA. Deixe vazio para usar a competência atual da planilha:',initialvalue=atual,parent=self)
        if valor is None:return
        valor=valor.strip()
        if valor:
            try:datetime.strptime(valor,'%m/%Y')
            except Exception:messagebox.showerror('Competência','Use o formato MM/AAAA.');return
        self.c['competencia']=valor;self.save_config_secure();self.consulta_cache=None
        self.refresh_dashboard();self.load(True)
        messagebox.showinfo('Competência','Competência selecionada: '+(valor or 'atual da planilha'))

    def alterar_fechamento(self):
        def worker():
            try:
                estado=self.api_get('fechamento');fechado=bool(estado.get('fechado'))
                pergunta='Reabrir a competência e permitir alterações?' if fechado else 'Fechar a competência? Novos pagamentos ficarão bloqueados até a reabertura.'
                def confirmar():
                    if not messagebox.askyesno('Fechamento mensal',pergunta):return
                    def executar():
                        try:
                            r=self.api_post({'acao':'reabrir_mes' if fechado else 'fechar_mes'})
                            self.after(0,lambda:messagebox.showinfo('Competência',r.get('mensagem') or ('Competência reaberta.' if fechado else 'Competência fechada.')))
                            self.after(0,self.refresh_dashboard)
                        except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Fechamento',str(e)))
                    threading.Thread(target=executar,daemon=True).start()
                self.after(0,confirmar)
            except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Fechamento',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def avancar_competencia(self):
        atual=self.c.get('competencia') or (self.consulta_cache or {}).get('competencia','')
        nova=simpledialog.askstring('Avançar competência','Nova competência MM/AAAA (vazio = próximo mês automático):',initialvalue='',parent=self)
        if nova is None:return
        nova=nova.strip()
        if nova:
            try:datetime.strptime(nova,'%m/%Y')
            except Exception:messagebox.showerror('Competência','Use o formato MM/AAAA.');return
        if not messagebox.askyesno('Confirmar avanço',f'Encerrar {atual or "a competência atual"} e avançar para {nova or "o próximo mês"}?'):return
        def worker():
            try:
                r=self.api_post({'acao':'avancar_competencia','competencia':nova})
                self.c['competencia']=r.get('competencia',nova);self.save_config_secure();self.consulta_cache=None
                self.after(0,self.refresh_dashboard);self.after(0,lambda:messagebox.showinfo('Competência',r.get('mensagem','Competência avançada.')))
            except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Competência',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def painel_anual(self):
        ano=simpledialog.askinteger('Painel anual','Ano:',initialvalue=datetime.now().year,parent=self)
        if not ano:return
        def abrir(r):
            w=tk.Toplevel(self);w.title('Painel anual');w.geometry('920x650');w.transient(self);ttk.Label(w,text='Arrecadação anual — '+str(ano),font=('Segoe UI',19,'bold'),padding=14).pack(anchor='w')
            graf=tk.Canvas(w,height=190,bg='white',highlightbackground='#D8E1EB',highlightthickness=1);graf.pack(fill='x',padx=14,pady=(0,10))
            meses=r.get('meses',[])
            def num(v):
                try:return float(str(v).replace('R$','').replace('.','').replace(',','.').strip())
                except:return 0
            maior=max([num(x.get('previsto')) for x in meses]+[1])
            for i,x in enumerate(meses):
                x0=28+i*72; base=165; hp=125*num(x.get('pago'))/maior; he=125*num(x.get('pendente'))/maior
                graf.create_rectangle(x0,base-hp,x0+22,base,fill='#1E8E5A',outline='');graf.create_rectangle(x0+24,base-he,x0+46,base,fill='#E98B19',outline='');graf.create_text(x0+23,178,text=str(i+1).zfill(2),fill='#44546A')
            graf.create_text(15,10,text='■ Recebido',fill='#1E8E5A',anchor='w');graf.create_text(110,10,text='■ Pendente',fill='#E98B19',anchor='w')
            tree=ttk.Treeview(w,columns=('mes','previsto','pago','pendente','qp','qd'),show='headings');
            for c,h,wd in [('mes','Mês',90),('previsto','Previsto',130),('pago','Recebido',130),('pendente','Pendente',130),('qp','Pagos',80),('qd','A pagar',80)]:tree.heading(c,text=h);tree.column(c,width=wd,anchor='center')
            for x in meses:tree.insert('','end',values=(x.get('competencia'),x.get('previsto'),x.get('pago'),x.get('pendente'),x.get('qtd_pagos'),x.get('qtd_pendentes')))
            tree.pack(fill='both',expand=True,padx=14,pady=(0,14))
        threading.Thread(target=lambda:self._thread_ui('anual',abrir,'Painel anual',ano=ano),daemon=True).start()

    def backup_online(self):
        if not messagebox.askyesno('Backup','Criar agora uma cópia completa da planilha no Google Drive?'):return
        def worker():
            try:r=self.api_post({'acao':'backup'});self.after(0,lambda:messagebox.showinfo('Backup','Cópia criada: '+r.get('nome','')))
            except Exception as e:self.after(0,lambda e=e:messagebox.showerror('Backup',str(e)))
        threading.Thread(target=worker,daemon=True).start()

    def _history_window(self,item,rows):
        w=tk.Toplevel(self); w.title('Histórico — '+str(item.get('morador',''))); w.geometry('920x500'); w.transient(self)
        ttk.Label(w,text=item.get('morador',''),font=('Segoe UI',18,'bold'),padding=14).pack(anchor='w')
        cols=('data','competencia','tipo','canal','resultado','obs'); heads=('Data/hora','Competência','Tipo','Canal','Resultado','Observações'); widths=(135,100,120,90,110,280)
        tree=ttk.Treeview(w,columns=cols,show='headings');
        for c,h,wd in zip(cols,heads,widths): tree.heading(c,text=h); tree.column(c,width=wd,anchor='w')
        for r in rows: tree.insert('','end',values=(r.get('data',''),r.get('competencia',''),r.get('tipo',''),r.get('canal',''),r.get('resultado',''),r.get('observacao','')))
        tree.pack(fill='both',expand=True,padx=14,pady=(0,14))

    def export_pdf(self,tipo,itens,resumo):
        if SimpleDocTemplate is None:
            messagebox.showerror('PDF','Instale os componentes executando novamente INICIAR_V8_2.bat.'); return
        destino=filedialog.asksaveasfilename(defaultextension='.pdf',filetypes=[('PDF','*.pdf')],initialfile=f"relacao_{tipo}_{resumo.get('competencia','').replace('/','-')}.pdf")
        if not destino: return
        styles=getSampleStyleSheet(); story=[Paragraph('Associação de Moradores',styles['Title']),Paragraph('Relatório '+tipo+' — '+resumo.get('competencia',''),styles['Heading2']),Spacer(1,12)]
        if tipo=='resumo':
            dados=[['Indicador','Valor'],['Moradores',resumo.get('total_moradores',0)],['Pagos',f"{resumo.get('qtd_pagos',0)} — {resumo.get('total_pago','')}"],['A pagar',f"{resumo.get('qtd_pendentes',0)} — {resumo.get('total_pendente','')}"],['Previsto',resumo.get('total_previsto','')]]
        elif tipo=='pagos':
            dados=[['Cód.','Morador','Valor','Data','Forma']]+[[x.get('codigo',''),x.get('morador',''),x.get('valor_pago',''),x.get('data_pagamento',''),x.get('forma_pagamento','')] for x in itens]
        else:
            dados=[['Cód.','Morador','Saldo','Vencimento']]+[[x.get('codigo',''),x.get('morador',''),x.get('saldo',''),x.get('vencimento','')] for x in itens]
        table=Table(dados,repeatRows=1); table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#123B66')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('FONTNAME',(0,0),(-1,0),'Helvetica-Bold'),('GRID',(0,0),(-1,-1),0.4,colors.HexColor('#CBD5E1')),('ROWBACKGROUNDS',(0,1),(-1,-1),[colors.white,colors.HexColor('#F2F6FA')]),('FONTSIZE',(0,0),(-1,-1),9),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('PADDING',(0,0),(-1,-1),6)])); story.append(table)
        SimpleDocTemplate(destino,pagesize=A4,rightMargin=28,leftMargin=28,topMargin=28,bottomMargin=28).build(story)
        messagebox.showinfo('PDF','Relatório PDF gerado com sucesso.')

    def export_excel(self,tipo,itens,resumo):
        if Workbook is None:
            messagebox.showerror('Excel','Instale os componentes executando novamente INICIAR_V8_2.bat.'); return
        destino=filedialog.asksaveasfilename(defaultextension='.xlsx',filetypes=[('Excel','*.xlsx')],initialfile=f"relacao_{tipo}_{resumo.get('competencia','').replace('/','-')}.xlsx")
        if not destino: return
        wb=Workbook(); ws=wb.active; ws.title='Relatório'
        if tipo=='resumo': rows=[['Indicador','Valor'],['Competência',resumo.get('competencia','')],['Moradores',resumo.get('total_moradores',0)],['Pagos',resumo.get('qtd_pagos',0)],['Total pago',resumo.get('total_pago','')],['A pagar',resumo.get('qtd_pendentes',0)],['Total pendente',resumo.get('total_pendente','')],['Previsto',resumo.get('total_previsto','')]]
        elif tipo=='pagos': rows=[['Código','Morador','Telefone','Valor pago','Data','Forma']]+[[x.get('codigo',''),x.get('morador',''),x.get('telefone',''),x.get('valor_pago',''),x.get('data_pagamento',''),x.get('forma_pagamento','')] for x in itens]
        else: rows=[['Código','Morador','Telefone','Saldo','Vencimento']]+[[x.get('codigo',''),x.get('morador',''),x.get('telefone',''),x.get('saldo',''),x.get('vencimento','')] for x in itens]
        for row in rows: ws.append(row)
        for cell in ws[1]: cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='123B66'); cell.alignment=Alignment(horizontal='center')
        for col in ws.columns: ws.column_dimensions[col[0].column_letter].width=min(max(len(str(c.value or '')) for c in col)+3,45)
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions; wb.save(destino)
        messagebox.showinfo('Excel','Relatório Excel gerado com sucesso.')

    def load(self,quiet=True):
        try:
            self.consulta_cache = None
            r=self.api_get('fila'); self.fila=r.get('fila',[]); self.i=0; self.retry=[]; self.persist(); self.showitem(); self.health.config(text='Conexão: OK — Planilha e Apps Script online',style='HealthOK.TLabel')
            self.summary.config(text=f"Fila: {len(self.fila)} | Vencidas: {r.get('vencidas',0)} | Lembretes: {r.get('lembretes',0)} | Total em aberto: {r.get('total_aberto','R$ 0,00')}")
            self.write(f'Fila atualizada: {len(self.fila)} cobrança(s).'); self.refresh_dashboard(); return len(self.fila)
        except Exception as e:
            self.health.config(text='Conexão: ERRO — '+str(e),style='HealthError.TLabel'); self.write('Erro: '+str(e))
            if not quiet: messagebox.showerror('Erro',str(e))
            return -1

    def persist(self):
        save_json(STATE, {'queue':self.fila,'index':self.i,'retry':self.retry})

    def offer_resume(self):
        if messagebox.askyesno('Retomar lote',f'Existe um lote interrompido no item {self.i+1} de {len(self.fila)}. Deseja retomar?'):
            self.show(); self.showitem()
        else:
            self.fila=[]; self.i=0; self.retry=[]; self.persist()

    def showitem(self):
        if not self.fila or self.i>=len(self.fila):
            self.nome.config(text='Fila concluída'); self.info.config(text=''); self.msg.delete('1.0','end'); self.prog.config(text=f'{len(self.fila)} de {len(self.fila)}'); return
        x=self.fila[self.i]; self.nome.config(text=x.get('morador','')); self.info.config(text=f"Prioridade: {x.get('prioridade','')} | {x.get('acao','')} | {x.get('competencia','')} | {x.get('saldo','')} | {x.get('telefone','')}")
        self.msg.delete('1.0','end'); self.msg.insert('1.0',x.get('mensagem','')); self.prog.config(text=f'{self.i+1} de {len(self.fila)}')

    def wa_url(self):
        x=self.fila[self.i]; n=''.join(c for c in x.get('telefone','') if c.isdigit()); n=n if n.startswith('55') else '55'+n
        return 'https://wa.me/'+n+'?text='+urllib.parse.quote(self.msg.get('1.0','end').strip())

    def openwa(self):
        if self.fila and self.i<len(self.fila): webbrowser.open(self.wa_url())

    def finish(self,result):
        if not self.fila or self.i>=len(self.fila): return
        x=self.fila[self.i]
        try:
            self.api_post({'acao':'marcar','linha':x['linha'],'resultado':result,'observacao':'Aplicativo v8.2.1'})
            self.consulta_cache = None
            self.write(f"{x['morador']}: {result}")
            if result=='ERRO': self.retry.append(x)
            self.i+=1; self.persist(); self.showitem()
        except Exception as e: messagebox.showerror('Erro',str(e))

    def batch(self):
        if self.running:return
        if not self.fila or self.i>=len(self.fila):
            if self.load(False)<=0:return
        self.running=True; self.stop=False; threading.Thread(target=self.worker,daemon=True).start()

    def worker(self):
        while self.running and self.i<len(self.fila) and not self.stop:
            x=self.fila[self.i]; self.after(0,lambda:webbrowser.open(self.wa_url())); time.sleep(max(3,self.c['delay']))
            if self.c['modo']=='AUTOMATICO' and pyautogui:
                try:
                    pyautogui.press('enter'); time.sleep(1)
                    self.api_post({'acao':'marcar','linha':x['linha'],'resultado':'ENVIADO','observacao':'Automático v8.2.1'})
                    self.consulta_cache = None
                    self.write(f"{x['morador']}: ENVIADO automático")
                    self.i+=1; self.persist(); self.after(0,self.showitem); continue
                except Exception as e: self.write('Erro automático: '+str(e))
            ev=threading.Event(); holder={'result':None}
            self.after(0,lambda ev=ev,holder=holder,x=x:self.confirm_popup(ev,holder,x))
            ev.wait()
            if holder['result']=='PARAR': break
            result=holder['result']
            try:
                self.api_post({'acao':'marcar','linha':x['linha'],'resultado':result,'observacao':'Confirmado no v8.2.1'})
                self.consulta_cache = None
                if result=='ERRO': self.retry.append(x)
                self.write(f"{x['morador']}: {result}"); self.i+=1; self.persist(); self.after(0,self.showitem)
            except Exception as e: self.write('Erro registro: '+str(e)); break
        if self.i>=len(self.fila) and self.retry:
            self.write(f'Retentativa pendente: {len(self.retry)} item(ns).')
            self.fila=self.retry.copy(); self.retry=[]; self.i=0; self.persist(); self.after(0,lambda:messagebox.showinfo('Retentativas','Os contatos com erro foram movidos para uma nova fila de retentativa.'))
        self.running=False

    def confirm_popup(self,ev,holder,x):
        w=tk.Toplevel(self); w.title('Confirmar envio'); w.geometry('430x250'); w.attributes('-topmost',True)
        ttk.Label(w,text=x.get('morador',''),font=('Segoe UI',15,'bold')).pack(pady=(20,8))
        ttk.Label(w,text='Confirme somente depois de enviar no WhatsApp.').pack()
        ttk.Label(w,text='Enter = Enviado | D = Adiar | S = Pular | E = Erro',font=('Segoe UI',10)).pack(pady=8)
        def done(r): holder['result']=r; w.destroy(); ev.set()
        b=ttk.Frame(w); b.pack(pady=16)
        ttk.Button(b,text='ENVIADO',command=lambda:done('ENVIADO')).pack(side='left',padx=4)
        ttk.Button(b,text='ADIAR 1 DIA',command=lambda:done('ADIAR_1_DIA')).pack(side='left',padx=4)
        ttk.Button(b,text='PULAR',command=lambda:done('PULADO')).pack(side='left',padx=4)
        ttk.Button(b,text='ERRO',command=lambda:done('ERRO')).pack(side='left',padx=4)
        w.bind('<Return>',lambda e:done('ENVIADO')); w.bind('d',lambda e:done('ADIAR_1_DIA')); w.bind('s',lambda e:done('PULADO')); w.bind('e',lambda e:done('ERRO'))
        w.protocol('WM_DELETE_WINDOW',lambda:done('PARAR')); w.focus_force()

    def monitor(self):
        while True:
            if not self.c.get('web_app_url') or not self.c.get('chave'):
                time.sleep(10)
                continue
            try:
                r=self.api_get('status'); self.health.after(0,lambda:self.health.config(text='Conexão: OK — Planilha e Apps Script online',style='HealthOK.TLabel'))
                qtd=r.get('qtd',0); sig=(qtd,r.get('assinatura',''))
                self.summary.after(0,lambda r=r:self.summary.config(text=f"Hoje: {r.get('qtd',0)} cobrança(s) | Vencidas: {r.get('vencidas',0)} | Lembretes: {r.get('lembretes',0)} | Em aberto: {r.get('total_aberto','R$ 0,00')}"))
                if qtd>0 and sig!=self.lastsig:
                    self.lastsig=sig; self.after(0,lambda q=qtd:self.alert(q))
                if qtd==0:self.lastsig=None
            except Exception as e:
                self.health.after(0,lambda e=e:self.health.config(text='Conexão: ERRO — '+str(e),style='HealthError.TLabel'))
            time.sleep(max(5,self.c['intervalo'])*60)

    def alert(self,qtd):
        self.show(); w=tk.Toplevel(self); w.title('Cobranças pendentes'); w.geometry('470x240'); w.attributes('-topmost',True)
        ttk.Label(w,text=f'Há {qtd} cobrança(s) para tratar.',font=('Segoe UI',16,'bold')).pack(pady=(28,10))
        ttk.Label(w,text='Clique para carregar a fila priorizada.').pack()
        ttk.Button(w,text='ABRIR FILA',command=lambda:(w.destroy(),self.load(False))).pack(pady=20)
        try:w.bell()
        except:pass

    def make_tray(self):
        if not pystray:return
        img=Image.new('RGB',(64,64),'white'); d=ImageDraw.Draw(img); d.rectangle((12,12,52,52),outline='black',width=4); d.line((20,26,44,26),fill='black',width=4); d.line((20,38,42,38),fill='black',width=4)
        menu=pystray.Menu(pystray.MenuItem('Abrir',lambda i,x:self.after(0,self.show)),pystray.MenuItem('Verificar agora',lambda i,x:self.after(0,lambda:self.load(True))),pystray.MenuItem('Sair',lambda i,x:self.after(0,self.quitall)))
        self.tray=pystray.Icon(APP,img,'Cobranças da Associação',menu); threading.Thread(target=self.tray.run,daemon=True).start()

    def write(self,s):
        line=time.strftime('[%Y-%m-%d %H:%M:%S] ')+s+'\n'
        try:
            with open(LOG_FILE,'a',encoding='utf-8') as f: f.write(line)
        except Exception: pass
        def x(): self.log.config(state='normal'); self.log.insert('end',time.strftime('[%H:%M:%S] ')+s+'\n'); self.log.see('end'); self.log.config(state='disabled')
        self.after(0,x)
    def show(self): self.deiconify(); self.lift()
    def hide(self): self.withdraw()
    def quitall(self):
        try:self.tray.stop()
        except:pass
        self.destroy()

if __name__ == '__main__':
    if not acquire_single_instance():
        root=tk.Tk(); root.withdraw()
        messagebox.showinfo('Cobranças da Associação','O Cobrador 8.2.1 já está aberto. Procure o ícone próximo ao relógio do Windows.')
        root.destroy()
    else:
        app=App()
        if getattr(app,'authenticated',False): app.mainloop()
